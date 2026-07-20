"""
1,2주차 작업 파일.

test_files/ 폴더를 재귀적으로 훑으면서, 지원하는 파일 형식(txt/csv/docx/pdf/xlsx)에서
텍스트를 추출하고, patterns.py의 탐지 함수들을 적용한다.
"""

from pathlib import Path
import docx
import pypdf
import openpyxl

from src.patterns import PII_DETECTORS

SUPPORTED_EXTENSIONS = {".txt", ".csv", ".docx", ".pdf", ".xlsx"}


def extract_text_from_file(filepath: Path) -> str:
    """파일 확장자에 따라 적절한 방식으로 텍스트를 추출."""
    suffix = filepath.suffix.lower()

    if suffix in (".txt", ".csv"):
        return filepath.read_text(encoding="utf-8", errors="ignore")

    elif suffix == ".docx":
        doc = docx.Document(filepath)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    elif suffix == ".pdf":
        reader = pypdf.PdfReader(filepath)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif suffix == ".xlsx":
        workbook = openpyxl.load_workbook(filepath, data_only=True)
        text_parts = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text_parts.append(str(cell))
        return "\n".join(text_parts)

    else:
        raise ValueError(f"지원하지 않는 파일 형식: {suffix}")


def scan_folder(folder: Path) -> dict:
    """
    folder 이하 모든 지원 파일을 재귀적으로 스캔해서
    파일별 탐지 결과를 딕셔너리로 반환.
    """
    results = {}

    for filepath in folder.rglob("*"):
        if not filepath.is_file():
            continue
        if filepath.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = extract_text_from_file(filepath)
        except Exception as e:
            print(f"[경고] {filepath} 읽기 실패: {e}")
            continue

        file_detections = {}
        for pii_type, detector_function in PII_DETECTORS.items():
            matches = detector_function(text)
            if matches:
                file_detections[pii_type] = matches

        if file_detections:
            results[str(filepath)] = file_detections

    return results
def scan_folder_with_llm_verification(folder: Path) -> dict:
    """
    scan_folder()와 동일하게 정규식으로 후보를 찾은 다음,
    각 후보를 LLM(verify_candidate)로 재검증해서 오탐으로 판단되면 결과에서 제외한다.
    무료 티어 속도 제한(rate limit)을 피하기 위해 호출 사이에 딜레이를 두고,
    실패 시 잠깐 기다렸다가 재시도한다.
    """
    import time
    from src.llm_verifier import verify_candidate

    raw_results = scan_folder(folder)
    verified_results = {}

    total_checked = 0
    total_failed = 0

    for filepath_str, detections in raw_results.items():
        filepath = Path(filepath_str)
        text = extract_text_from_file(filepath)

        verified_detections = {}
        for pii_type, matches in detections.items():
            kept_matches = []
            for value, start, end in matches:
                context = text  # 더미 파일이 짧음. So 전체를 문맥으로 사용

                total_checked += 1
                is_valid = None

                # 최대 3번까지 재시도
                for attempt in range(3):
                    try:
                        is_valid = verify_candidate(context, value, pii_type)
                        break
                    except Exception as e:
                        wait_time = 5 * (attempt + 1)  # 5초,10초,15초 점점 늘려가며 대기
                        print(f"  [경고] LLM 호출 실패 ({value}), {wait_time}초 대기 후 재시도... ({e})")
                        time.sleep(wait_time)

                if is_valid is None:
                    print(f"  [실패] {value}: 3번 재시도 후에도 실패 → 정규식 결과 그대로 유지")
                    is_valid = True
                    total_failed += 1

                status = "유지" if is_valid else "제외(오탐 판단)"
                print(f"  [{total_checked}] {pii_type}: {value} → {status}")

                if is_valid:
                    kept_matches.append((value, start, end))

                time.sleep(3)  # 다음 호출 전에 3초 대기(속도 제한 방지)

            if kept_matches:
                verified_detections[pii_type] = kept_matches

        if verified_detections:
            verified_results[filepath_str] = verified_detections

    print(f"\n총 {total_checked}건 검증, 그중 {total_failed}건은 API 실패로 정규식 결과 유지")
    return verified_results

if __name__ == "__main__":
    results = scan_folder(Path("test_files"))
    print(results)